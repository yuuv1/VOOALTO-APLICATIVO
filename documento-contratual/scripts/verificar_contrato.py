#!/usr/bin/env python3
"""Valida o modelo contratual em PDF e Word (em branco e com ficha técnica + imagens)."""
from __future__ import annotations

import argparse
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader, PdfWriter
import pymupdf

from gerar_contrato import (
    validate_data, money, amount_words, build_docx, ContractPDF,
    add_digital_signature_fields, register_fonts, ROOT, OUT_NAME,
    SCOPE_TEXT, ACCEPTANCE_TEXT, M,
)
from extrair_identidade import SOURCE, SOURCE_SHA256

EXPECTED_TEXT_FIELDS = {
    "cliente_nome_razao_social",
    "cliente_cpf_cnpj",
    "cliente_telefone",
    "cliente_email",
    "cliente_endereco",
    "cliente_cidade_uf_cep",
    "item_01_especificacao", "item_01_grade", "item_01_qtd", "item_01_unitario", "item_01_subtotal",
    "item_02_especificacao", "item_02_grade", "item_02_qtd", "item_02_unitario", "item_02_subtotal",
    "item_03_especificacao", "item_03_grade", "item_03_qtd", "item_03_unitario", "item_03_subtotal",
    "item_04_especificacao", "item_04_grade", "item_04_qtd", "item_04_unitario", "item_04_subtotal",
    "producao_quantidade_total",
    "producao_valor_total",
    "producao_prazo_entrega",
    "producao_observacoes",
    "pagamento_entrada_valor",
    "pagamento_saldo_valor",
    "pagamento_forma",
    "local_data_assinatura",
}
EXPECTED_SIGNATURE_FIELDS = {"assinatura_digital_contratante"}


def normalize(text: str) -> str:
    return " ".join(text.split())


def word_text(doc) -> str:
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(p.text for p in cell.paragraphs)
    for sec in doc.sections:
        chunks.extend(p.text for p in sec.footer.paragraphs)
    return normalize(" ".join(chunks))


class ContractTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = json.loads((ROOT / "dados/contrato.json").read_text(encoding="utf-8"))
        cls.pdf_path = ROOT / "entregas" / f"{OUT_NAME}.pdf"
        cls.docx_path = ROOT / "entregas" / f"{OUT_NAME}.docx"
        cls.reader = PdfReader(cls.pdf_path)
        cls.fields = cls.reader.get_fields()
        cls.text = normalize("\n".join(page.extract_text() for page in cls.reader.pages))
        cls.doc = Document(cls.docx_path)
        cls.doc_text = word_text(cls.doc)

    def test_original_preservado(self):
        self.assertEqual(sha256(SOURCE.read_bytes()).hexdigest(), SOURCE_SHA256)

    def test_pdf_uma_folha_a4(self):
        self.assertEqual(len(self.reader.pages), 1)
        page = self.reader.pages[0]
        self.assertAlmostEqual(float(page.mediabox.width), 595.2756, places=2)
        self.assertAlmostEqual(float(page.mediabox.height), 841.8898, places=2)

    def test_empresa_campos_cliente_producao_e_assinaturas_em_ambos(self):
        for text in [self.text, self.doc_text]:
            for expected in [
                "VOOALTO UNIFORMES",
                "51.919.138/0001-20",
                "(62) 98135-6948",
                "Av. Jacinto Alves de Abreu, 212",
                "74370-661",
                "NOME COMPLETO OU RAZÃO SOCIAL",
                "CPF / CNPJ",
                "TELEFONE / WHATSAPP",
                "E-MAIL",
                "ENDEREÇO COMPLETO",
                "CIDADE / UF",
                "DADOS DA PRODUÇÃO E ESPECIFICAÇÕES DO PEDIDO",
                "MODELO / TECIDO / COR / PERSONALIZAÇÃO",
                "GRADE / TAM.",
                "CROQUIS / IMAGENS DE REFERÊNCIA DAS CAMISETAS (FICHA TÉCNICA)",
                "CONDIÇÕES DE PAGAMENTO",
                "ENTRADA / SINAL NA APROVAÇÃO",
                "SALDO NA ENTREGA FINAL",
                "FORMA DE PAGAMENTO",
                "LOCAL E DATA",
                SCOPE_TEXT,
                ACCEPTANCE_TEXT,
                "CONTRATADA (EMPRESA)",
                "CONTRATANTE (CLIENTE)",
                "Esio Andrade — Proprietário",
                "Assinatura do(a) Contratante / Representante",
            ]:
                self.assertIn(expected, text)
            lower = text.lower()
            self.assertNotIn("evaluation", lower)
            self.assertNotIn("aspose", lower)
            self.assertNotIn("financeiro", lower)
            # Garante que nenhum dado antigo de pedido foi deixado preenchido por engano
            self.assertNotIn("4.312,80", text)
            self.assertNotIn("2.156,40", text)
            self.assertNotIn("faixa lateral", lower)

    def test_todos_os_campos_vazios_e_sem_assinatura_criptografica(self):
        self.assertEqual(set(self.fields), EXPECTED_TEXT_FIELDS | EXPECTED_SIGNATURE_FIELDS)
        for name in EXPECTED_TEXT_FIELDS:
            self.assertEqual(self.fields[name]["/FT"], "/Tx")
            self.assertEqual(self.fields[name].get("/V", ""), "")
        for name in EXPECTED_SIGNATURE_FIELDS:
            self.assertEqual(self.fields[name]["/FT"], "/Sig")
            self.assertNotIn("/V", self.fields[name])
        self.assertNotIn(b"/ByteRange", self.pdf_path.read_bytes())
        self.assertNotIn("/Perms", self.reader.trailer["/Root"])
        self.assertFalse(self.reader.is_encrypted)

    def test_pdf_preenchivel_sem_alterar_original(self):
        writer = PdfWriter()
        writer.clone_document_from_reader(self.reader)
        example = "EMPRESA EXEMPLO LTDA"
        writer.update_page_form_field_values(
            writer.pages[0],
            {"cliente_nome_razao_social": example, "cliente_cpf_cnpj": "12.345.678/0001-99"},
            auto_regenerate=False,
        )
        output = BytesIO()
        writer.write(output)
        output.seek(0)
        filled = PdfReader(output)
        self.assertEqual(filled.get_fields()["cliente_nome_razao_social"]["/V"], example)
        self.assertEqual(filled.get_fields()["cliente_cpf_cnpj"]["/V"], "12.345.678/0001-99")
        self.assertNotIn("/V", filled.get_fields()["assinatura_digital_contratante"])
        self.assertEqual(PdfReader(self.pdf_path).get_fields()["cliente_nome_razao_social"]["/V"], "")

    def test_pdf_campos_e_textos_dentro_da_folha(self):
        with pymupdf.open(self.pdf_path) as doc:
            page = doc[0]
            for widget in page.widgets():
                self.assertGreaterEqual(widget.rect.x0, M - 0.5)
                self.assertLessEqual(widget.rect.x1, page.rect.width - M + 0.5)
                self.assertGreaterEqual(widget.rect.y0, 100)
                self.assertLessEqual(widget.rect.y1, 800)
            for block in page.get_text("dict", flags=0)["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        x0, y0, x1, y1 = span["bbox"]
                        self.assertGreaterEqual(x0, M - 1)
                        self.assertLessEqual(x1, page.rect.width - M + 1)
                        self.assertGreaterEqual(y0, 0)
                        self.assertLessEqual(y1, page.rect.height)

    def test_word_a4_sem_quebras_forcadas(self):
        self.assertEqual(len(self.doc.sections), 1)
        section = self.doc.sections[0]
        self.assertAlmostEqual(section.page_width.cm, 21.0, places=2)
        self.assertAlmostEqual(section.page_height.cm, 29.7, places=2)
        self.assertAlmostEqual(section.left_margin.cm, 1.5, places=2)
        self.assertAlmostEqual(section.right_margin.cm, 1.5, places=2)
        with ZipFile(self.docx_path) as archive:
            xml = archive.read("word/document.xml").decode()
            self.assertNotIn('w:type="page"', xml)
            self.assertNotIn('w:hRule="exact"', xml)

    def test_word_larguras_das_celulas_consistentes_sem_recorte(self):
        section = self.doc.sections[0]
        available = (section.page_width - section.left_margin - section.right_margin) // 635
        for table in self.doc.tables:
            pr = table._tbl.tblPr
            self.assertEqual(pr.find(qn("w:tblLayout")).get(qn("w:type")), "fixed")
            self.assertEqual(pr.find(qn("w:jc")).get(qn("w:val")), "left")
            grid = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
            width = int(pr.find(qn("w:tblW")).get(qn("w:w")))
            self.assertEqual(width, sum(grid))
            self.assertLessEqual(width, available + 2)
            for row in table.rows:
                cell_widths = [int(tc.get_or_add_tcPr().find(qn("w:tcW")).get(qn("w:w"))) for tc in row._tr.tc_lst]
                self.assertEqual(cell_widths, grid)
                self.assertEqual(sum(cell_widths), width)
                self.assertIsNotNone(row._tr.get_or_add_trPr().find(qn("w:cantSplit")))

    def test_word_imagens_originais_sem_limite_de_cabecalho(self):
        self.assertEqual(len(self.doc.inline_shapes), 2)
        for header in [self.doc.sections[0].header, self.doc.sections[0].first_page_header]:
            self.assertEqual(len(header.tables), 0)
        with ZipFile(self.docx_path) as archive:
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
            self.assertEqual(len(media), 2)
        for shape in self.doc.inline_shapes:
            self.assertGreater(shape.width, 0)
            self.assertGreater(shape.height, 0)
            self.assertLess(shape.width.cm, 6)

    def test_pdf_logo_assinatura_e_fontes_incorporadas(self):
        with pymupdf.open(self.pdf_path) as doc:
            self.assertGreaterEqual(len(doc[0].get_images()), 2)
        fonts = self.reader.pages[0]["/Resources"]["/Font"]
        embedded = 0
        for ref in fonts.values():
            font = ref.get_object()
            if "DejaVu" in str(font.get("/BaseFont", "")):
                self.assertIn("/FontFile2", font["/FontDescriptor"].get_object())
                embedded += 1
        self.assertGreaterEqual(embedded, 2)

    def test_modo_ficha_tecnica_preenchida_com_imagens_e_validacao(self):
        sample = json.loads(json.dumps(self.data))
        sample["producao"]["itens"] = [
            {
                "numero": "01",
                "modelo": "Camiseta gola polo institucional",
                "tecido": "Malha fria PV antipilling",
                "cor": "Azul marinho",
                "personalizacao": "Silk digital peito e costas",
                "grade": "P: 4 · M: 10 · G: 6",
                "quantidade": 20,
                "valor_unitario": "64.90",
                "subtotal": "1298.00",
            }
        ]
        sample["producao"]["imagens_referencia"] = [
            {"titulo": "Ref. 01 — Polo Frente", "arquivo": "assets/logo-vooalto.png"},
        ]
        sample["pagamento"]["entrada_percentual"] = 50
        sample["pagamento"]["saldo_percentual"] = 50
        comp = validate_data(sample)
        self.assertTrue(comp["has_items"])
        self.assertEqual(comp["total_display"], money("1298.00"))
        self.assertEqual(comp["advance_display"], money("649.00"))
        self.assertEqual(comp["words_display"], amount_words("1298.00"))

        with tempfile.TemporaryDirectory() as tmp:
            tmp_dir = Path(tmp)
            register_fonts()
            base_pdf = tmp_dir / "base.pdf"
            out_pdf = tmp_dir / "out.pdf"
            out_docx = tmp_dir / "out.docx"
            builder = ContractPDF(base_pdf, sample, comp)
            builder.build()
            add_digital_signature_fields(base_pdf, out_pdf, builder.signatures)
            build_docx(sample, out_docx, comp)
            r = PdfReader(out_pdf)
            d = Document(out_docx)
            self.assertIn("R$ 1.298,00", normalize(r.pages[0].extract_text()))
            self.assertIn("R$ 1.298,00", word_text(d))
            self.assertEqual(len(d.inline_shapes), 3)

        # Divergência de subtotal deve ser rejeitada
        broken = json.loads(json.dumps(sample))
        broken["producao"]["itens"][0]["subtotal"] = "1299.00"
        with self.assertRaises(ValueError):
            validate_data(broken)


class MPTOContractTests(unittest.TestCase):
    MPTO_DIR = ROOT / "contratos" / "MPTO_Ministerio_Publico_do_Tocantins"

    @classmethod
    def setUpClass(cls):
        cls.pdf_path = cls.MPTO_DIR / "Contrato_MPTO_Ficha_3374.pdf"
        cls.docx_path = cls.MPTO_DIR / "Contrato_MPTO_Ficha_3374.docx"
        cls.json_path = cls.MPTO_DIR / "dados_contrato_mpto.json"
        cls.ficha_path = cls.MPTO_DIR / "FICHA_TECNICA_3374_MPTO_Original.pdf"
        cls.data = json.loads(cls.json_path.read_text(encoding="utf-8"))
        cls.reader = PdfReader(cls.pdf_path)
        cls.fields = cls.reader.get_fields()
        cls.text = normalize("\n".join(page.extract_text() for page in cls.reader.pages))
        cls.doc = Document(cls.docx_path)
        cls.doc_text = word_text(cls.doc)

    def test_pasta_separada_mpto_completa(self):
        self.assertTrue(self.pdf_path.is_file())
        self.assertTrue(self.docx_path.is_file())
        self.assertTrue(self.json_path.is_file())
        self.assertTrue(self.ficha_path.is_file())
        self.assertEqual(len(list((self.MPTO_DIR / "referencias").glob("ref_0*.jpg"))), 5)

    def test_mpto_pdf_uma_folha_a4_sem_transbordo(self):
        self.assertEqual(len(self.reader.pages), 1)
        page = self.reader.pages[0]
        self.assertAlmostEqual(float(page.mediabox.width), 595.2756, places=2)
        self.assertAlmostEqual(float(page.mediabox.height), 841.8898, places=2)
        with pymupdf.open(self.pdf_path) as doc:
            p0 = doc[0]
            for widget in p0.widgets():
                self.assertGreaterEqual(widget.rect.x0, M - 0.5)
                self.assertLessEqual(widget.rect.x1, p0.rect.width - M + 0.5)
                self.assertGreaterEqual(widget.rect.y0, 80)
                self.assertLessEqual(widget.rect.y1, 825)
            for block in p0.get_text("dict", flags=0)["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        x0, y0, x1, y1 = span["bbox"]
                        self.assertGreaterEqual(x0, M - 1)
                        self.assertLessEqual(x1, p0.rect.width - M + 1)
                        self.assertGreaterEqual(y0, 0)
                        self.assertLessEqual(y1, p0.rect.height)

    def test_mpto_campos_preenchidos_em_pdf_e_word(self):
        expected_fields = {
            "cliente_nome_razao_social": "Ministério Público do Estado do Tocantins (MPTO / Procuradoria-Geral de Justiça)",
            "cliente_cpf_cnpj": "01.786.078/0001-46",
            "cliente_telefone": "(63) 99994-1987",
            "cliente_endereco": "Quadra 202 Norte, Conjunto 01, Lotes 05/06, Avenida LO-04, Plano Diretor Norte",
            "cliente_cidade_uf_cep": "Palmas - TO · CEP 77.006-218",
        }
        for fname, val in expected_fields.items():
            self.assertEqual(self.fields[fname]["/V"], val)
            self.assertIn(val, self.doc_text)
        self.assertIn("30 dias", self.fields["producao_prazo_entrega"]["/V"])
        self.assertIn("3374", self.fields["producao_observacoes"]["/V"])
        self.assertIn("22 de outubro de 2026", self.fields["local_data_assinatura"]["/V"])
        self.assertIn("30 dias", self.doc_text)
        self.assertIn("3374", self.doc_text)
        self.assertIn("22 de outubro de 2026", self.doc_text)

    def test_mpto_itens_quantidades_valores_e_parcelas_sincronizados(self):
        for txt in [self.text, self.doc_text]:
            for expected in [
                "Camiseta gola polo manga curta",
                "Camiseta gola redonda manga curta",
                "99 PEÇAS",
                "R$ 1.852,80",
                "R$ 926,40",
                "R$ 1.563,30",
                "R$ 598,80",
                "R$ 5.540,10",
                "R$ 2.770,05",
                "Esio Andrade — Proprietário",
                "Assinatura do(a) Contratante / Representante",
            ]:
                self.assertIn(expected, txt)

    def test_mpto_assinatura_contratante_em_branco_e_7_imagens(self):
        sig_field = self.fields["assinatura_digital_contratante"]
        self.assertEqual(sig_field["/FT"], "/Sig")
        self.assertNotIn("/V", sig_field)
        self.assertNotIn(b"/ByteRange", self.pdf_path.read_bytes())
        self.assertNotIn("/Perms", self.reader.trailer["/Root"])
        with pymupdf.open(self.pdf_path) as doc:
            # 1 logo + 5 camisetas + 1 assinatura Esio = 7 imagens (0 assinatura cliente)
            self.assertEqual(len(doc[0].get_images()), 7)
        self.assertEqual(len(self.doc.inline_shapes), 7)

    def test_mpto_word_a4_tabelas_alinhadas_sem_quebras(self):
        section = self.doc.sections[0]
        available = (section.page_width - section.left_margin - section.right_margin) // 635
        with ZipFile(self.docx_path) as archive:
            xml = archive.read("word/document.xml").decode()
            self.assertNotIn('w:type="page"', xml)
        for table in self.doc.tables:
            pr = table._tbl.tblPr
            grid = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
            width = int(pr.find(qn("w:tblW")).get(qn("w:w")))
            self.assertEqual(width, sum(grid))
            self.assertLessEqual(width, available + 2)


class PWAServerWindowsTests(unittest.TestCase):
    def test_servidor_funciona_sem_console_pythonw_e_mime_js_correto(self):
        import sys
        import threading
        import urllib.request
        import servidor_pwa

        old_out, old_err = sys.stdout, sys.stderr
        srv = servidor_pwa.VooaltoHTTPServer(("127.0.0.1", 0), servidor_pwa.VooaltoPWAHandler)
        port = srv.server_address[1]
        t = threading.Thread(target=srv.serve_forever, daemon=True)
        t.start()
        try:
            sys.stdout = None
            sys.stderr = None

            with urllib.request.urlopen(f"http://127.0.0.1:{port}/api/status", timeout=3) as r_status:
                self.assertEqual(r_status.status, 200)
                payload = json.loads(r_status.read().decode("utf-8"))
                self.assertTrue(payload["ok"])
                self.assertEqual(payload["versao"], servidor_pwa.SERVER_VERSION)

            with urllib.request.urlopen(f"http://127.0.0.1:{port}/", timeout=3) as r_html:
                self.assertEqual(r_html.status, 200)
                self.assertIn("text/html", r_html.headers.get("Content-Type", ""))
                self.assertIn('id="root"', r_html.read().decode("utf-8"))

            js_file = next((ROOT / "docs/assets").glob("index-*.js"))
            mjs_file = next((ROOT / "docs/assets").glob("pdf.worker.min-*.mjs"))
            for asset in [js_file, mjs_file]:
                with urllib.request.urlopen(f"http://127.0.0.1:{port}/assets/{asset.name}", timeout=3) as r_js:
                    self.assertEqual(r_js.status, 200)
                    self.assertIn("text/javascript", r_js.headers.get("Content-Type", ""))
                    self.assertGreater(len(r_js.read()), 1000)
        finally:
            sys.stdout = old_out
            sys.stderr = old_err
            srv.shutdown()
            srv.server_close()



def verify_word_render(path: Path) -> None:
    """Confere uma renderização local do DOCX em PDF; nunca a promove a arquivo de entrega."""
    with pymupdf.open(path) as doc:
        assert len(doc) == 1, f"O Word renderizado deve caber em uma única folha A4 (gerou {len(doc)})."
        page = doc[0]
        assert abs(page.rect.width - 595.28) < 1 and abs(page.rect.height - 841.89) < 1, "A folha deve ser A4."
        text = normalize(page.get_text(sort=True))
        for expected in [
            "VOOALTO UNIFORMES",
            "51.919.138/0001-20",
            "(62) 98135-6948",
            "NOME COMPLETO OU RAZÃO SOCIAL",
            "CPF / CNPJ",
            "TELEFONE / WHATSAPP",
            "E-MAIL",
            "ENDEREÇO COMPLETO",
            "DADOS DA PRODUÇÃO E ESPECIFICAÇÕES DO PEDIDO",
            "CROQUIS / IMAGENS DE REFERÊNCIA DAS CAMISETAS (FICHA TÉCNICA)",
            "CONDIÇÕES DE PAGAMENTO",
            "Esio Andrade — Proprietário",
            "Assinatura do(a) Contratante / Representante",
        ]:
            assert expected in text, f"Conteúdo ausente ou recortado no Word: {expected}"
        for block in page.get_text("dict", flags=0)["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    span_txt = span["text"]
                    if "Aspose" in span_txt or "Evaluation" in span_txt or "products.aspose" in span_txt:
                        continue
                    x0, y0, x1, y1 = span["bbox"]
                    assert 30 <= x0 and x1 <= page.rect.width - 25, f"Recorte lateral no Word: {span_txt}"
                    assert 0 <= y0 and y1 <= page.rect.height, f"Recorte vertical no Word: {span_txt}"
        assert len(page.get_images()) >= 2, "Logo e assinatura precisam estar na mesma folha."
    print("Renderização local do Word validada: 1 folha A4, todos os campos, logo e assinaturas sem recortes.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--word-renderizado", type=Path)
    args, remaining = parser.parse_known_args()
    result = unittest.main(argv=["verificar_contrato.py", *remaining], verbosity=2, exit=False).result
    if not result.wasSuccessful():
        raise SystemExit(1)
    if args.word_renderizado:
        verify_word_render(args.word_renderizado)
